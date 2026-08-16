"""文档格式解析测试:PDF/Word/Excel/PowerPoint 文本提取。"""
import io

import pytest

from app.rag.documents import extract_text


def test_extract_pdf():
    import pymupdf

    buf = io.BytesIO()
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 72), "PDF 资产清单示例内容 2024", fontname="china-s")
    doc.save(buf)
    doc.close()
    text = extract_text(buf.getvalue(), "sample.pdf")
    assert "PDF 资产清单" in text


def test_extract_docx():
    from docx import Document

    doc = Document()
    doc.add_paragraph("Word 文档标题")
    doc.add_paragraph("这是正文段落。")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "资产名称"
    table.rows[0].cells[1].text = "数量"
    buf = io.BytesIO()
    doc.save(buf)
    text = extract_text(buf.getvalue(), "sample.docx")
    assert "Word 文档标题" in text
    assert "资产名称" in text
    assert "数量" in text


def test_extract_xlsx():
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "资产明细"
    ws.append(["资产", "数量", "单价"])
    ws.append(["电脑", 10, 5000.5])
    buf = io.BytesIO()
    wb.save(buf)
    text = extract_text(buf.getvalue(), "sample.xlsx")
    assert "资产明细" in text
    assert "电脑" in text
    assert "5000.5" in text


def test_extract_xls():
    xlwt = pytest.importorskip("xlwt")
    wb = xlwt.Workbook()
    ws = wb.add_sheet("资产明细")
    ws.write(0, 0, "资产")
    ws.write(0, 1, "数量")
    ws.write(1, 0, "电脑")
    ws.write(1, 1, 10)
    buf = io.BytesIO()
    wb.save(buf)
    text = extract_text(buf.getvalue(), "sample.xls")
    assert "资产明细" in text
    assert "电脑" in text


def test_extract_pptx():
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    box.text_frame.text = "季度汇报 PPT 内容"
    buf = io.BytesIO()
    prs.save(buf)
    text = extract_text(buf.getvalue(), "sample.pptx")
    assert "季度汇报" in text


def test_unsupported_format_returns_hint():
    text = extract_text(b"\x89PNG\r\n\x1a\n\x00\x00", "image.png")
    assert "未支持的格式" in text
